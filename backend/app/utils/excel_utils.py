import shutil
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .template_source import resolve_template_path

SHEET_NAME = "入力用"
ROW_OFFSET = 22
YEAR_CELL = "J7"
MONTH_CELL = "J8"

CATEGORY1_LABELS = {
    "1": "1：出勤",
    "2": "2：欠勤",
    "3": "3：有給休暇(全日)",
    "4": "4：有給休暇(前半)",
    "5": "5：休日出勤",
    "6": "6：法定休日出勤",
    "7": "7：特別休暇",
    "8": "8：休業",
    "9": "9：有給休暇（後半）",
}

CATEGORY2_LABELS = {
    "1": "1：遅刻",
    "2": "2：早退",
    "3": "3：遅刻＋早退",
    "4": "4：休業遅刻",
    "5": "5：休業早退",
    "6": "6：休業遅刻＋休業早退",
}

COLUMN_MAP = {
    "category1": 11,  # K
    "category2": 15,  # O
    "start_time": 18,  # R
    "end_time": 20,  # T
    "break_minutes": 22,  # V
    "notes": 32,  # AF
}


def _build_output_path(year: int, month: int, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    filename = f"{month}月_業務進捗報告_島.xlsx"
    return output_dir / filename


def _copy_template(template_path: Path, output_path: Path) -> None:
    if not template_path.exists():
        raise FileNotFoundError(f"Template file not found: {template_path}")
    shutil.copy(template_path, output_path)


def _format_break_minutes(minutes: int) -> str:
    return f"{minutes // 60}:{minutes % 60:02d}"


def _write_attendance_row(sheet, attendance, month: int) -> None:
    if not attendance.date or attendance.date.month != month:
        return

    row = attendance.date.day + ROW_OFFSET
    sheet.cell(
        row=row,
        column=COLUMN_MAP["category1"],
        value=CATEGORY1_LABELS.get(str(attendance.category1), attendance.category1),
    )
    sheet.cell(
        row=row,
        column=COLUMN_MAP["category2"],
        value=CATEGORY2_LABELS.get(str(attendance.category2), attendance.category2),
    )
    sheet.cell(
        row=row,
        column=COLUMN_MAP["start_time"],
        value=attendance.start_time.strftime("%H:%M") if attendance.start_time else None,
    )
    sheet.cell(
        row=row,
        column=COLUMN_MAP["end_time"],
        value=attendance.end_time.strftime("%H:%M") if attendance.end_time else None,
    )
    sheet.cell(
        row=row,
        column=COLUMN_MAP["break_minutes"],
        value=_format_break_minutes(attendance.break_minutes),
    )
    if attendance.notes:
        sheet.cell(row=row, column=COLUMN_MAP["notes"], value=attendance.notes)


def export_attendance_to_excel(
    attendances: Iterable,
    year: int,
    month: int,
    template_path: str,
    output_dir: str,
) -> str:
    template_file = resolve_template_path(template_path)
    output_path = _build_output_path(year, month, Path(output_dir))
    _copy_template(template_file, output_path)

    workbook = load_workbook(output_path)
    if SHEET_NAME in workbook.sheetnames:
        sheet = workbook[SHEET_NAME]
    else:
        sheet = workbook.active

    # 管理年・管理月をテンプレートにセット
    sheet[YEAR_CELL] = year
    sheet[MONTH_CELL] = month

    for attendance in attendances:
        _write_attendance_row(sheet, attendance, month)

    workbook.save(output_path)
    return str(output_path)
